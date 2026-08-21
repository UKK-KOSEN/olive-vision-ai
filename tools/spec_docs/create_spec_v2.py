"""OliveVision AI 技術仕様書 v2 - コードベース完全一致版
src/runtime.py 全2,222行を精査し、コード上の値・関数・ロジックと
一切のズレがないよう作成している。"""

import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ── スタイル ──
style = doc.styles["Normal"]
style.font.name = "Yu Gothic"
style.font.size = Pt(10.5)
for lv in range(1, 4):
    hs = doc.styles[f"Heading {lv}"]
    hs.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)
    hs.font.bold = True
    hs.font.size = Pt(18 - lv * 2)

GREEN = RGBColor(0x1B, 0x5E, 0x20)

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def tbl(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True; r.font.size = Pt(9)
    for rd in rows:
        row = t.add_row()
        for i, v in enumerate(rd):
            c = row.cells[i]
            c.text = str(v)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

def title_page(main, sub, info):
    doc.add_paragraph(); doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(main); r.font.size = Pt(28); r.bold = True; r.font.color.rgb = GREEN
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub); r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x61,0x61,0x61)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(info); r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x9E,0x9E,0x9E)
    doc.add_page_break()

# タイトル
title_page(
    "OliveVision AI 技術仕様書",
    "AI搭載 オリーブ樹長期モニタリングシステム\n"
    "システムアーキテクチャ / 検出アルゴリズム / パラメータ仕様",
    "バージョン: 2.0\n最終更新: 2026年8月17日\n"
    "対象ファイル: src/runtime.py (2,222行) / config/runtime.yaml / olivevision.py"
)

# ── 各セクションを別モジュールから読み込み ──
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from spec_sections import build_all
build_all(doc, code, tbl)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "OliveVision_AI_技術仕様書_v2.docx")
doc.save(out)
print(f"saved: {out}")
